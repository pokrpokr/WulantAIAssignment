import click
import fitz
import cv2
import os
import numpy as np
import traceback
from functools import reduce
from openpyxl import Workbook

# .select([0, 1]) select PDF 1st & 2nd page

def read_pdf(file_path, page_start, page_end):
    # fitz.open() can not be regconised by vscode, using .Document() instead
    doc = fitz.Document(file_path)
    
    if page_start == 0 and page_start == page_end:
        p_start, p_end = 0, len(doc) 
    elif page_start > 0 and page_start < len(doc) and page_start == page_end:
        p_start, p_end = page_start-1, page_end
    elif page_start > 0 and page_start < len(doc) and page_end > page_start and page_end < len(doc):
        p_start, p_end = page_start-1, page_end
    else:
        raise Exception("Can not find pages in PDF")
    
    # doc.select([x for x in range(p_start, p_end)])
    paragraph_blocks = {}
    for page_index in range(p_start, p_end):
        page = doc[page_index]
              
        # blocks = page.get_text("blocks")

        # # Initialize variables to keep track of columns
        # columns = []
        # prev_x = None

        # # Iterate through each text block
        # for block in blocks:
        #     if block[6] == 0:
        #         x0, y0, x1, y1, line = block[:5]  # Unpack the text and rectangle coordinates

        #         # Check if this text block belongs to the same column
        #         if prev_x is not None and x0 > prev_x:
        #             columns[-1].append(line)  # Add text to the last column
        #         else:
        #             columns.append([line])  # Start a new column
        #         prev_x = x0

        # # Combine text blocks within each column into paragraphs
        # paragraphs = ['\n'.join(column) for column in columns]
      
        paragraphs = []
        text_blocks = []
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            # skip image
            if block["type"] == 1:
                continue
            else:
                text_blocks.append(block)
        
        text_blocks.sort(key=lambda x: (x["bbox"][0], x["bbox"][1]))
        
        for tb in text_blocks:
            lines = []
            for line in tb["lines"]:
                span_text = [sp["text"] for sp in line["spans"]]
                lines.append(reduce(lambda x, y: x+y, span_text))
            paragraphs.append(reduce(lambda x, y: x+"\n"+y, lines))
        paragraph_blocks[page_index+1] = paragraphs
    doc.close()
    return paragraph_blocks
        
def export_to_excel(data: dict, excel_name):
    wb = Workbook()
    wb.remove(wb.active)
    for page, page_data in data.items():
        ws = wb.create_sheet(title=f"Page{page}")
        for row in page_data:
            ws.append([row])
    wb.save(excel_name)

# OCR
from pytesseract import image_to_string

def pdf_to_image(file_path, page_start, page_end, zoom=2):
    # document = fitz.open(pdf_path)
    doc = fitz.Document(file_path)
    # text_blocks = []
    
    if page_start == 0 and page_start == page_end:
        p_start, p_end = 0, len(doc) 
    elif page_start > 0 and page_start < len(doc) and page_start == page_end:
        p_start, p_end = page_start-1, page_end
    elif page_start > 0 and page_start < len(doc) and page_end > page_start:
        p_start, p_end = page_start-1, page_end
    else:
        raise Exception("Can not find pages in PDF")
    
    for page_index in range(p_start, p_end):
        page = doc[page_index]
        mat = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
        img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
        yield img

def drawImgContours(windows_name, image, contours):
    # image: page_img
    # index: -1 find all contours
    # color:
    # line_width:
    temp = image.copy()
    cv2.drawContours(temp,contours, -1, (0,255,0), 1)
    cv2.imshow(windows_name, temp)
    key = cv2.waitKey(0)  

    if key == 27:  # type ESC
        cv2.destroyAllWindows()
    elif key == ord('q'):  # type q
        cv2.destroyAllWindows()

def find_text_blocks(image):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    gray = cv2.GaussianBlur(gray,(3,3), 1)

    thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    
    contours, _ = cv2.findContours(thresh, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    drawImgContours("test", image, contours)
    print(f"Find contours {len(contours)}")
    text_blocks = []
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        if w * h > 25:  
            text_blocks.append((x, y, w, h))
    return text_blocks

def ocr_text_blocks(image, text_blocks):
    text_blocks.sort(key=lambda b: (b[1], b[0]))
    texts = []
    for block in text_blocks:
        x, y, w, h = block
        text_roi = image[y:y+h, x:x+w]
        text = image_to_string(text_roi, lang='chi_sim+eng')
        texts.append(text.strip())
    return texts


@click.command()
@click.option("--src_pdf", "-sp", type=str, prompt="Please enter PDF file name")
@click.option("--expt_name", "-en", type=str, prompt="Please enter output excel file name")
@click.option("--ver", "-ver", type=int, default=1, prompt="ver: 1 without using opencv; ver: 2 using opencv")
@click.option("--pg_start", "-pgs", type=int, default=0, prompt="PDF start page")
@click.option("--pg_end", "-pge", type=int, default=0, prompt="PDF end page")
def main(src_pdf, expt_name, ver, pg_start, pg_end):
    if ver == 1:
        print("----- Processing PDF -----")
        result = read_pdf(os.path.join(os.getcwd()+"/test_pdfs/", src_pdf), page_start=pg_start, page_end=pg_end)
        export_to_excel(result, os.path.join(os.getcwd()+"/export_excels/", expt_name))
        print(f"----- File location: {os.path.join(os.getcwd()+'/export_excels/', expt_name)} -----")
        print("----- Export Finished -----")
    elif ver == 2:
        print("----- Processing PDF with opencv -----")
        for page_image in pdf_to_image(os.path.join(os.getcwd()+"/test_pdfs/", src_pdf), page_start=pg_start, page_end=pg_end):
            blocks = find_text_blocks(page_image)
            texts = ocr_text_blocks(page_image, blocks)
            print(texts)
            print("-------------------")
            # for text in texts:
            #     print(text)
        print("----- Export Finished -----")
    else:
        raise Exception("Not other versions")

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"Error occur! Msg: {str(e)}")
        print(traceback.format_exc())